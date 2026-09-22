import os
from typing import Literal

import openpyxl as pyxl
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table as ExcelTable
from openpyxl.worksheet.table import TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from Excel.Chart import Chart
from Excel.CollapsibleTable import CollapsibleTable
from Excel.Table import Table
from helpers import is_file_open

# --- central place to expand styles ---
FONT_TOKENS = {
    "bold": {"bold": True},
    "underline": {"underline": "single"},
    "orange": {"color": "FFA500"},
    "red": {"color": "FF0000"},
    "blue": {"color": "0000FF"},
}

FILL_TOKENS = {
    "redbg": "F25757",
    "greenbg": "57F257",
    "yellowbg": "F2F257",
}

ALIGN_TOKENS = {"left": "left", "center": "center", "right": "right"}

NUMFMT_TOKENS = {
    "num": "#,##0",
    "twodecimal": "#,##0.00",
    "percent": "0%",
    "date": "m/d/yyyy",
    "datetime": "m/d/yyyy h:mm:ss AM/PM",
    "time": "h:mm:ss AM/PM",
}

BORDER_TOKENS = {"border": Border(bottom=Side(style="thin"))}


class Workbook:
    """A class representing the Excel workbook to be modified
    """
    def __init__(self, file_name: str, mode: Literal["new", "existing"]):
        """instantiates the workbook object (the Excel file): new or existing

        Args:
            file_name (str): the name of the new or existing file
            mode (Literal["new", "existing"]): the type of file being accessed

        Raises:
            ValueError: invalid mode passed
            FileNotFoundError: existing file is not found
            PermissionError: existing file is open
        """
        # validate file mode
        if mode not in ("new", "existing"):
            raise ValueError(f"mode must be 'new' or 'existing', got {mode!r}")

        self.file_name = file_name
        if mode == 'new':
            self.wb = pyxl.Workbook()
            del self.wb['Sheet']
        elif mode == 'existing':
            if not os.path.exists(self.file_name):
                raise FileNotFoundError(f"No such file: {self.file_name!r}")
            elif is_file_open(self.file_name):
                raise PermissionError("The file is open in Excel. Close it before proceeding.")
            else:
                self.wb = pyxl.load_workbook(self.file_name)


    def new_sheet(self, sheet_name: str, position: int | None= None) -> Worksheet:
        """Create a new worksheet in the Excel workbook

        Args:
            sheet_name (str): the name of the sheet
            position (int): an optional position of the sheet in the tabs

        Returns:
            Worksheet: the openpyxl worksheet object
        """
        sheet = self.wb.create_sheet(sheet_name, position)
        self.wb.active = self.wb.worksheets.index(sheet)
        return sheet


    def make_sheet_active(self, sheet: Worksheet):
        """makes a sheet active for editing

        Args:
            sheet (Worksheet): the sheet to make active
        """
        self.wb.active = self.wb.worksheets.index(sheet)
        return sheet


    def get_active_sheet(self, sheet: Worksheet | None= None) -> Worksheet:
        """gets the active sheet, sets if different from currently active sheet

        Args:
            sheet (Worksheet | None, optional): the sheet object. Defaults to None.

        Raises:
            ValueError: no active sheets

        Returns:
            Worksheet: the active sheet
        """
        if sheet is not None:
            ws = self.make_sheet_active(sheet)
        else:
            if self.wb.active is not None:
                ws = self.wb.active
            else:
                raise ValueError('Workbook has no active sheet')

        return ws


    def get_active_sheetname(self) -> str:
        """Gets the name of the active sheet

        Returns:
            str: the name of the active sheet
        """
        return self.wb.active.title


    def get_sheets(self) -> list[Worksheet]:
        """Gets all openpyxl worksheet objects

        Returns:
            list[Worksheet]: all openpyxl worksheet objects
        """
        return self.wb.worksheets


    def hide_sheet(self):
        ...


    def extend_column_widths(self, widths: list[int | float], start_col: int = 1, sheet: Worksheet | None = None):
        """extends the column widths of the specified sheet or active if not specified

        Args:
            widths (list[int | float]): the widths of the corresponding columns
            start_col (int, optional): the column to start extending from. Defaults to 1.
            sheet (Worksheet | None, optional): the sheet to apply extensions to. Defaults to None.
        """
        sheet = self.get_active_sheet(sheet)
        for offset, width in enumerate(widths):
            col_letter = get_column_letter(start_col + offset)
            sheet.column_dimensions[col_letter].width = width


    def add_table(self, table: Table, sheet: Worksheet | None= None):
        """writes and formats the table in native Excel

        Args:
            table (Table): the table with all values and formatting assigned
            sheet (Worksheet | None, optional): the sheet to add the table to. Defaults to None.
        """
        ws = self.get_active_sheet(sheet)

        column_formats = table.get_column_formats()
        total_formats = table.get_total_formats()
        header_formats = table.get_header_formats()
        total_operations = table.get_total_operations()

        # determine offset based on total row
        if total_operations != {} and table.total_row == 'top':
            table_total_row = table.header_row
            table.header_row += 1
            table.data_row += 1
        else:
            table_total_row = table.data_row + len(table.data)

        # write the title if one is specified
        if table.title is not None:
            last_col = table.start_col + len(table.data.columns) - 1
            ws.merge_cells(start_row=table.title_row, start_column=table.start_col,
                        end_row=table.title_row, end_column=last_col)
            title_cell = ws.cell(row=table.title_row, column=table.start_col, value=table.title)
            title_cell.alignment = Alignment(horizontal=table.title_align)

        # loop through and write the header
        for col_idx, col_name in enumerate(table.data.columns, start=table.start_col):
            self.write_value(ws, table.header_row, col_idx, col_name, col_name, header_formats)

        # loop through and write the data
        for row_idx, row in enumerate(table.data.itertuples(index=False), start=table.data_row):
            for col_idx, (col_name, value) in enumerate(zip(table.data.columns, row), start=table.start_col):
                self.write_value(ws, row_idx, col_idx, value, col_name, column_formats)

        # loop though and write the total row
        for col_idx, col_name in enumerate(table.data.columns, start=table.start_col):
            if total_operations.get(col_name) is not None:
                self.write_value(ws, table_total_row, col_idx, total_operations[col_name], col_name, total_formats)

        # add the actual Excel Table object, referencing the range you just filled
        last_col_letter = get_column_letter(table.start_col + len(table.data.columns) - 1)
        first_col_letter = get_column_letter(table.start_col)
        last_data_row = table.data_row + len(table.data) - 1

        table_ref = f"{first_col_letter}{table.header_row}:{last_col_letter}{last_data_row}"

        excel_table = ExcelTable(displayName=table.name, ref=table_ref)
        excel_table.tableStyleInfo = TableStyleInfo(
            name=table.style,
            showRowStripes=True
        )
        ws.add_table(excel_table)


    def add_collapsible_table(self, table: CollapsibleTable, sheet: Worksheet | None= None):
        ws = self.get_active_sheet(sheet)


    def add_chart(self, chart: Chart, sheet: Worksheet | None= None):
        ws = self.get_active_sheet(sheet)


    def close(self):
        """Safely saves and closes the openpyxl Excel file
        """
        self.wb.save(self.file_name)
        self.wb.close()




    def __build_named_style(self, format_str: str) -> NamedStyle:
        """build a named style openpyxl object for formatting a cell

        Args:
            format_str (str): a '_' seperated string of formats to apply

        Raises:
            ValueError: Unknown format token

        Returns:
            NamedStyle: the combined cell format
        """
        # define default formatting
        font_kwargs, fill_color, align, number_format, border = {}, None, None, "General", None

        # loop through each format presented and add per token type
        for token in format_str.lower().split("_"):
            if token in FONT_TOKENS:
                font_kwargs.update(FONT_TOKENS[token])
            elif token in FILL_TOKENS:
                fill_color = FILL_TOKENS[token]
            elif token in ALIGN_TOKENS:
                align = ALIGN_TOKENS[token]
            elif token in NUMFMT_TOKENS:
                number_format = NUMFMT_TOKENS[token]
            elif token in BORDER_TOKENS:
                border = BORDER_TOKENS[token]
            else:
                raise ValueError(f"Unknown format token: {token!r} in {format_str!r}")

        # build the named style
        style = NamedStyle(name=format_str)
        if font_kwargs:
            style.font = Font(**font_kwargs)
        if fill_color:
            style.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        if align:
            style.alignment = Alignment(horizontal=align)
        if border:
            style.border = border
        style.number_format = number_format
        return style


    def set_cell(self, sheet: Worksheet, row: int, col: int, value, format_str: str | None = None):
        """wrapper for setting a cell value with formatting

        Args:
            sheet (Worksheet): the sheet object of the sheet being worked on
            row (int): the row index of the cell to apply to
            col (int): the column index of the cell to apply to
            value (any): the value to set the cell to
            format_str (str): the format to apply to the cell
        """
        # set cell value
        cell = sheet.cell(row=row, column=col, value=value)
        # add named style if not added already and not None or ''
        if format_str:
            if format_str not in self.wb.named_styles:
                self.wb.add_named_style(self.__build_named_style(format_str))
                # set format
            cell.style = format_str


    def write_value(self, sheet: Worksheet, row: int, col: int, value, col_name: str, formats: dict):
        """generic format-based write

        Args:
            sheet (Worksheet): worksheet object to work on
            row (int): row index of the cell to modify
            col (int): col index of the cell to modify
            value (any): the value to put in the cell
            col_name (str): the name of column
            formats (dict): the formats and conditions to apply to a value
        """
        # check if the column has a format
        if col_name in formats:
            # loop through each format
            formatted_check = False # checks if any formats were applied
            for format_str, condition in formats[col_name]:
                # if there is a condition
                if condition is not None and eval(condition.replace(col_name, str(value))):
                    self.set_cell(sheet, row, col, value, format_str)
                    formatted_check = True
                    break
                # no condition
                elif condition is None:
                    self.set_cell(sheet, row, col, value, format_str)
            # cell met no conditions
            if formatted_check is False:
                self.set_cell(sheet, row, col, value)
        # no format or condition
        else:
            self.set_cell(sheet, row, col, value)